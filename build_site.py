# -*- coding: utf-8 -*-
"""Статический сайт «Маляр в Сануре» из книги SEO. Выход: docs/ (GitHub Pages)."""
import json, os, re, shutil, html
import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
CFG = json.load(open(os.path.join(ROOT, "site_config.json"), encoding="utf-8"))
BOOK = os.path.join(ROOT, "SEO-малярные-работы-Бали-2026-09-25.xlsx")
OUT = os.path.join(ROOT, "docs")
B = CFG["base_path"].rstrip("/")
E = html.escape

wb = openpyxl.load_workbook(BOOK, read_only=True)
pages = {r[0]: dict(id=r[0], url=r[1], parent=r[2], type=r[3], h1=r[5], title=r[6], desc=r[8])
         for r in wb["03 Структура"].iter_rows(min_row=2, values_only=True)}
sem = list(wb["09 Семантика"].iter_rows(min_row=2, values_only=True))
QW = ("как ", "чем ", "почему ", "что ", "можно ли ", "сколько ", "какой ", "какая ", "какую ", "нужно ли ", "когда ")

# Тексты услуг: только проверяемое, без цен, сроков-обещаний и выдуманных кейсов
TXT = {
 "S01": ("Крашу стены и потолки в виллах, квартирах и домах Санура: от одной комнаты до всего дома. Перед покраской проверяю основание на влагу и плесень — во влажном климате это главная причина, по которой свежая краска пузырится и отслаивается.",
         ["Осмотр и замер влажности стен", "Защита мебели и пола", "Очистка, обработка антисептиком при необходимости", "Шпаклёвка трещин и шлифовка", "Грунт и два слоя краски", "Уборка после работы"]),
 "S02": ("Покраска фасадов вилл у моря требует другого подхода: соль, солнце и дожди разрушают покрытие быстрее, чем в глубине острова. Подбираю фасадные краски под прибрежные условия и планирую работы на сухой сезон.",
         ["Осмотр трещин, отслоений и зелёного налёта", "Мойка и биоцидная обработка", "Ремонт трещин эластичными составами", "Грунт и фасадная краска", "Покраска цоколя, парапетов, металлических элементов"]),
 "S03": ("Не всегда нужно перекрашивать всю стену. Подкрашиваю следы после арендаторов, царапины, пятна после протечек и сколы — с подбором цвета, чтобы не было заметных пятен.",
         ["Подбор и колеровка цвета по образцу", "Локальная подготовка и грунт", "Подкраска с растушёвкой", "Честная рекомендация, если стену лучше перекрасить целиком"]),
 "S04": ("Плесень в Сануре — частая проблема: морская влажность и сезон дождей. Удаляю плесень со стен, потолков, дерева и швов, нахожу причину и защищаю поверхность, чтобы она не вернулась через месяц.",
         ["Поиск источника влаги", "Механическая очистка", "Обработка фунгицидом", "Просушка", "Антигрибковый грунт и краска", "Рекомендации по вентиляции"]),
 "S05": ("Антигрибковая обработка — профилактика и обязательный этап перед покраской поражённых поверхностей: стены, потолки, деревянные балки, шкафы, мебель.",
         ["Выбор состава под поверхность", "Нанесение и выдержка", "Повторная обработка при глубоком поражении", "Финишное покрытие"]),
 "S06": ("Покрываю лаком двери, мебель, столешницы, лестницы, полы и уличное дерево. Во влажном воздухе лак ведёт себя иначе: мутнеет, дольше сохнет, отслаивается при ошибках подготовки — поэтому работаю по погоде и с промежуточной шлифовкой.",
         ["Снятие старого покрытия", "Шлифовка", "Грунт/порозаполнитель", "Нанесение лака в несколько слоёв", "Промежуточная шлифовка", "Лак для улицы с UV-защитой"]),
 "S07": ("Масло и пропитка — для тика, террас, уличной мебели и дерева, которое должно дышать. Защищают от влаги, солнца и насекомых и легко обновляются без полной шлифовки.",
         ["Очистка и шлифовка", "Антисептическая пропитка", "Масло или масло-воск в 2–3 слоя", "План обновления покрытия"]),
 "S08": ("Восстанавливаю деревянную мебель и двери: снимаю старое покрытие, убираю плесень и пятна, заделываю дефекты и заново покрываю маслом или лаком.",
         ["Оценка состояния", "Снятие старого покрытия", "Ремонт сколов и трещин", "Тонировка при необходимости", "Финишное покрытие"]),
 "S09": ("Деревянная терраса у бассейна или сада сереет и темнеет от солнца, воды и плесени. Шлифую, очищаю и покрываю маслом для террас.",
         ["Мойка и отбеливание дерева", "Шлифовка", "Обработка от плесени", "Масло для террас", "Рекомендации по уходу"]),
 "S10": ("У моря металл ржавеет быстро. Крашу ворота, заборы, перила и решётки с полной подготовкой: удаление ржавчины, антикоррозийный грунт, эмаль.",
         ["Удаление ржавчины", "Преобразователь и антикоррозийный грунт", "Эмаль в 2 слоя", "Покраска сложных элементов кистью"]),
 "S11": ("Качество покраски определяется подготовкой. Заделываю трещины, ремонтирую штукатурку, шпаклюю и выравниваю стены перед покраской.",
         ["Расшивка и заделка трещин", "Ремонт отбитой штукатурки", "Шпаклёвка", "Шлифовка", "Грунтование"]),
 "S12": ("Влагозащитные покрытия и гидрофобизация — для стен, которые мокнут в дождь, цоколей, парапетов и камня. Важно: покрытие не заменяет устранение протечки, сначала ищем источник воды.",
         ["Диагностика источника влаги", "Гидрофобизатор для камня и бетона", "Влагозащитные краски", "Обработка швов и примыканий"]),
}

IMG = {"S01": "roller", "S02": "facade", "S03": "ladder", "S04": "bali", "S05": "roller", "S06": "varnish",
       "S07": "wood", "S08": "restore", "S09": "deck", "S10": "ladder", "S11": "hero", "S12": "facade"}
ALT = {"hero": "Мастер красит наружную стену дома", "facade": "Покраска белого фасада здания", "ladder": "Мастер на стремянке у фасада",
       "restore": "Мастер работает с деревом", "bali": "Строитель на объекте на Бали", "wood": "Балийский мастер работает с деревом",
       "roller": "Валик с краской на стене", "varnish": "Покрытие деревянной доски кистью", "villa": "Вилла с бассейном на Бали",
       "deck": "Деревянная терраса после дождя"}
WA_SVG = '<svg class="wa-ico" viewBox="0 0 24 24"><path d="M12 2a10 10 0 0 0-8.6 15.1L2 22l5-1.3A10 10 0 1 0 12 2zm5.3 14.1c-.2.6-1.3 1.2-1.8 1.2-.5.1-1 .2-3.3-.7-2.8-1.1-4.6-4-4.7-4.2-.1-.2-1.1-1.5-1.1-2.9s.7-2 1-2.3c.2-.3.5-.3.7-.3h.5c.2 0 .4 0 .6.5l.8 2c.1.2.1.4 0 .5l-.4.6-.4.4c-.1.2-.3.3-.1.6.2.3.8 1.3 1.7 2.1 1.2 1 2.1 1.4 2.4 1.5.3.1.5.1.6-.1l.9-1c.2-.3.4-.2.6-.1l1.9.9c.3.1.5.2.5.3.1.2.1.8-.2 1.3z"/></svg>'

def contact_html(small=False):
    sm = " btn-sm" if small else ""
    parts = []
    if CFG["whatsapp"]:
        parts.append(f'<a class="btn btn-wa{sm}" href="https://wa.me/{re.sub(r"[^0-9]", "", CFG["whatsapp"])}">{WA_SVG}WhatsApp</a>')
    if CFG["telegram"]:
        parts.append(f'<a class="btn btn-ghost{sm}" href="https://t.me/{CFG["telegram"].lstrip("@")}">Telegram</a>')
    if CFG["phone"]:
        parts.append(f'<a class="btn btn-ghost{sm}" href="tel:{re.sub(r"[^0-9+]", "", CFG["phone"])}">Позвонить</a>')
    return " ".join(parts) or f'<a class="btn btn-wa{sm}" href="{B}/kontakty/">{WA_SVG}Написать мастеру</a>'

def img(name, cls="", eager=False):
    load = 'fetchpriority="high"' if eager else 'loading="lazy"'
    return f'<img class="{cls}" src="{B}/img/{name}.webp" alt="{E(ALT[name])}" {load} decoding="async">'

CSS = """:root{--bg:#FBF8F3;--surface:#fff;--ink:#1B2B27;--muted:#5B6B66;--line:#E8E2D8;--brand:#1F5E4F;--brand-2:#E9F2EE;--accent:#F2A65A;--accent-2:#FCEBD8;--wa:#25D366;--radius:20px}
*{box-sizing:border-box}html{scroll-behavior:smooth}
body{margin:0;background:var(--bg);color:var(--ink);font:17px/1.6 Inter,system-ui,sans-serif;-webkit-font-smoothing:antialiased}
h1,h2,h3{font-family:Fraunces,Georgia,serif;font-weight:600;line-height:1.15;letter-spacing:-.01em;margin:0 0 .5em}
h1{font-size:clamp(2.1rem,5vw,3.5rem)}h2{font-size:clamp(1.6rem,3vw,2.3rem)}h3{font-size:1.2rem}
a{color:var(--brand)}p{margin:0 0 1em}img{max-width:100%;display:block}
.wrap{max-width:1120px;margin:0 auto;padding:0 20px}
header{position:sticky;top:0;z-index:20;background:rgba(251,248,243,.88);backdrop-filter:blur(12px);border-bottom:1px solid var(--line)}
.nav{display:flex;align-items:center;justify-content:space-between;height:68px;gap:16px}
.logo{font-family:Fraunces,serif;font-weight:600;font-size:1.3rem;color:var(--ink);text-decoration:none}.logo span{color:var(--accent)}
.links{display:flex;gap:26px}.links a{color:var(--muted);text-decoration:none;font-size:.95rem}.links a:hover{color:var(--ink)}
.btn{display:inline-flex;align-items:center;gap:10px;padding:14px 22px;border-radius:999px;font-weight:600;text-decoration:none;font-size:1rem;transition:transform .15s,box-shadow .15s}
.btn:hover{transform:translateY(-1px)}.btn-wa{background:var(--wa);color:#fff;box-shadow:0 6px 20px rgba(37,211,102,.3)}
.btn-ghost{background:transparent;color:var(--ink);border:1.5px solid var(--line)}.btn-sm{padding:9px 16px;font-size:.9rem}
.wa-ico{width:20px;height:20px;fill:currentColor;flex:none}
.hero{padding-top:56px;padding-bottom:48px;display:grid;grid-template-columns:1.05fr .95fr;gap:48px;align-items:center}
.eyebrow{display:inline-block;background:var(--brand-2);color:var(--brand);padding:6px 14px;border-radius:999px;font-size:.85rem;font-weight:600;margin-bottom:18px}
.lead{font-size:1.18rem;color:var(--muted);max-width:34em}
.cta-row{display:flex;gap:12px;flex-wrap:wrap;margin:26px 0 20px}
.trust{display:flex;gap:10px 22px;flex-wrap:wrap;color:var(--muted);font-size:.92rem;list-style:none;padding:0;margin:0}
.trust li::before{content:"✓ ";color:var(--brand);font-weight:700}
.photo{position:relative;border-radius:28px;overflow:hidden;aspect-ratio:4/5;background:var(--accent-2)}
.photo img{width:100%;height:100%;object-fit:cover}
.chip{position:absolute;left:18px;bottom:18px;background:#fff;border-radius:14px;padding:12px 16px;box-shadow:0 10px 30px rgba(0,0,0,.1);font-size:.9rem;max-width:80%}
.chip b{display:block;font-family:Fraunces,serif;font-size:1.05rem}
section{padding:60px 0}.alt-bg{background:var(--surface);border-top:1px solid var(--line);border-bottom:1px solid var(--line)}
.sec-head{max-width:660px;margin-bottom:34px}.sec-head p{color:var(--muted)}
.grid{display:grid;gap:18px;grid-template-columns:repeat(auto-fill,minmax(250px,1fr))}
.card{background:var(--surface);border:1px solid var(--line);border-radius:var(--radius);overflow:hidden;text-decoration:none;color:inherit;display:flex;flex-direction:column;transition:border-color .15s,box-shadow .15s,transform .15s}
a.card:hover{border-color:#cfd9d4;box-shadow:0 12px 30px rgba(27,43,39,.08);transform:translateY(-2px)}
.card img{aspect-ratio:16/10;object-fit:cover;width:100%}.card .body{padding:18px 20px 22px}
.card h3{margin-bottom:6px}.card p{color:var(--muted);font-size:.95rem;margin:0}
.split{display:grid;grid-template-columns:1fr 1fr;gap:48px;align-items:center}
.split .photo{aspect-ratio:4/3}
.steps{counter-reset:s;display:grid;grid-template-columns:repeat(4,1fr);gap:18px;padding:0;list-style:none}
.steps li{counter-increment:s;background:var(--surface);border-radius:var(--radius);padding:24px;border:1px solid var(--line)}
.steps li::before{content:counter(s);display:grid;place-items:center;width:36px;height:36px;border-radius:50%;background:var(--accent-2);color:#b0632a;font-weight:700;margin-bottom:12px}
.steps b{display:block;margin-bottom:4px}
.area{display:flex;flex-wrap:wrap;gap:10px;padding:0;list-style:none}.area li{background:var(--brand-2);color:var(--brand);border-radius:999px;padding:8px 16px;font-weight:600;font-size:.92rem}
.checks{list-style:none;padding:0;display:grid;gap:10px}.checks li{background:var(--surface);border:1px solid var(--line);border-radius:14px;padding:14px 18px}
.checks li::before{content:"✓";color:var(--brand);font-weight:700;margin-right:10px}
details{background:var(--surface);border:1px solid var(--line);border-radius:14px;padding:14px 18px;margin-bottom:10px}summary{cursor:pointer;font-weight:600}
details p{margin:.6em 0 0;color:var(--muted)}
.crumbs{font-size:.88rem;color:var(--muted);padding-top:20px}.crumbs a{color:var(--muted)}
.band{background:var(--brand);color:#fff;border-radius:28px;padding:44px;display:grid;grid-template-columns:1.4fr 1fr;gap:24px;align-items:center}
.band h2{color:#fff}.band p{color:#d5e6df;margin:0}.band .cta-row{margin:0;justify-content:flex-end}
.page-hero{display:grid;grid-template-columns:1.1fr .9fr;gap:40px;align-items:center;padding-top:28px;padding-bottom:20px}.page-hero .photo{aspect-ratio:4/3}
.credit{font-size:.78rem;color:var(--muted);margin-top:8px}
footer{border-top:1px solid var(--line);padding:36px 0 90px;color:var(--muted);font-size:.9rem}
.float{position:fixed;right:16px;bottom:16px;z-index:30}
@media (max-width:860px){.hero,.split,.page-hero,.band{grid-template-columns:1fr}.hero{padding-top:28px}.links{display:none}
.steps{grid-template-columns:1fr 1fr}.band{padding:30px}.band .cta-row{justify-content:flex-start}.photo{aspect-ratio:4/3}}
@media (max-width:520px){.steps{grid-template-columns:1fr}body{font-size:16px}}"""

def page(path, title, desc, body, crumbs=None, schema=None, og=None):
    canon = CFG["site_url"].rstrip("/") + path
    cr = ""
    if crumbs:
        cr = '<div class="wrap crumbs">' + " › ".join(f'<a href="{B}{u}">{E(t)}</a>' if u else E(t) for t, u in crumbs) + "</div>"
    sch = "".join(f'<script type="application/ld+json">{json.dumps(x, ensure_ascii=False)}</script>' for x in (schema or []))
    ogi = f'<meta property="og:image" content="{CFG["site_url"]}/img/{og or "hero"}.webp">'
    doc = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{E(title)}</title><meta name="description" content="{E(desc)}"><link rel="canonical" href="{canon}">
<meta property="og:title" content="{E(title)}"><meta property="og:description" content="{E(desc)}">{ogi}<meta name="theme-color" content="#FBF8F3">
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Fraunces:opsz,wght@9..144,600&family=Inter:wght@400;600;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="{B}/style.css">{sch}</head><body>
<header><div class="wrap nav"><a class="logo" href="{B}/">Маляр<span>·</span>Санур</a>
<nav class="links"><a href="{B}/uslugi/">Услуги</a><a href="{B}/ceny/">Цены</a><a href="{B}/raboty/">Работы</a><a href="{B}/kontakty/">Контакты</a></nav>
{contact_html(True).split("</a>")[0] + "</a>"}</div></header>
<main>{cr}{body}</main>
<footer><div class="wrap">{E(CFG["name"])} — покраска, удаление плесени и покрытие дерева в Сануре, Бали.<br>
Фотографии на сайте — иллюстрации (Pexels, Unsplash), не работы мастера. Фото объектов появятся в разделе «Работы».</div></footer>
</body></html>"""
    d = os.path.join(OUT, path.strip("/"))
    os.makedirs(d, exist_ok=True)
    open(os.path.join(d, "index.html"), "w", encoding="utf-8").write(doc)
    return canon

def faq_for(pid, n=5):
    qs = []
    for r in sem:
        if r[7] == pid and r[3] == "ru" and str(r[2]).startswith(QW) and "санур" not in r[2]:
            qs.append(r[2])
    return qs[:n]

biz = {"@context": "https://schema.org", "@type": "HousePainter", "name": CFG["name"], "url": CFG["site_url"] + "/",
       "image": CFG["site_url"] + "/img/hero.webp", "areaServed": {"@type": "Place", "name": "Sanur, Bali, Indonesia"}}
if CFG["phone"]:
    biz["telephone"] = CFG["phone"]
if CFG["gbp_url"]:
    biz["sameAs"] = [CFG["gbp_url"]]

if os.path.exists(OUT):
    shutil.rmtree(OUT)
os.makedirs(os.path.join(OUT, "img"))
for f in os.listdir(os.path.join(ROOT, "img_src")):
    if f.endswith(".webp"):
        shutil.copy(os.path.join(ROOT, "img_src", f), os.path.join(OUT, "img", f))
open(os.path.join(OUT, "style.css"), "w", encoding="utf-8").write(CSS)
open(os.path.join(OUT, ".nojekyll"), "w").write("")
urls = []
svcs = [p for p in pages.values() if p["type"] == "услуга"]

def short(t, n=110):
    return t if len(t) <= n else t[:n].rsplit(" ", 1)[0] + "…"

def svc_cards(items=None):
    return '<div class="grid">' + "".join(
        f'<a class="card" href="{B}{p["url"]}">{img(IMG[p["id"]])}<div class="body"><h3>{E(p["h1"].replace(" в Сануре", ""))}</h3><p>{E(short(TXT[p["id"]][0]))}</p></div></a>'
        for p in (items or svcs)) + "</div>"

AREA = ["Санур Кайя", "Санур Кауф", "Семаванг", "Синду", "Мертасари", "Батуджимбар"]
STEPS = [("Фото в WhatsApp", "Пришлите фото и примерную площадь."), ("Осмотр в Сануре", "Приезжаю, смотрю основание и влажность."),
         ("Смета", "Объём, материалы, сроки — до начала работ."), ("Работа и уборка", "Делаю, убираю, показываю результат.")]

home = pages["P000"]
urls.append(page("/", home["title"], home["desc"], f"""
<div class="wrap hero"><div>
<span class="eyebrow">Только Санур · выезд на осмотр</span>
<h1>Маляр в Сануре: покраска, плесень, лак для дерева</h1>
<p class="lead">Крашу стены и фасады вилл, убираю плесень и защищаю дерево лаком и маслом. Знаю, как влажность и морской воздух Санура разрушают покрытия, поэтому начинаю с причины, а не с закрашивания.</p>
<div class="cta-row">{contact_html()}<a class="btn btn-ghost" href="{B}/uslugi/">Все услуги</a></div>
<ul class="trust"><li>Расчёт по фото</li><li>Материалы под влажный климат</li><li>Уборка после работы</li></ul>
</div><div class="photo">{img("hero", eager=True)}<div class="chip"><b>Санур, Бали</b>Покраска · плесень · дерево</div></div></div>

<section class="alt-bg"><div class="wrap"><div class="sec-head"><h2>Услуги</h2><p>От подкраски одной стены до фасада виллы и террасы у бассейна.</p></div>
{svc_cards()}</div></section>

<section><div class="wrap split"><div class="photo">{img("facade")}</div><div>
<h2>Почему у моря покрытия служат меньше</h2>
<p>Санур стоит на побережье: соль в воздухе ускоряет коррозию металла, солнце разрушает лак на уличном дереве, а влажность и сезон дождей дают плесень на стенах и потолках.</p>
<ul class="checks"><li>Проверяю влажность стены до покраски</li><li>Обрабатываю плесень, а не закрашиваю её</li><li>Подбираю краски и лаки под прибрежный климат</li><li>Планирую фасады на сухой сезон</li></ul>
</div></div></section>

<section class="alt-bg"><div class="wrap"><div class="sec-head"><h2>Как проходит работа</h2></div>
<ol class="steps">{"".join(f"<li><b>{E(a)}</b>{E(b)}</li>" for a, b in STEPS)}</ol></div></section>

<section><div class="wrap split"><div>
<h2>Дерево: лак, масло, реставрация</h2>
<p>Тик, двери, мебель, террасы у бассейна. Во влажном воздухе лак мутнеет и отслаивается при ошибках подготовки — работаю по погоде, с шлифовкой между слоями.</p>
<div class="cta-row"><a class="btn btn-ghost" href="{B}/uslugi/lakirovka-dereva/">Покрытие лаком</a><a class="btn btn-ghost" href="{B}/uslugi/terrasa-dekking/">Терраса</a></div>
</div><div class="photo">{img("wood")}</div></div></section>

<section class="alt-bg"><div class="wrap"><div class="sec-head"><h2>Где работаю</h2><p>Только Санур — поэтому быстро приезжаю на осмотр и на гарантийные вопросы.</p></div>
<ul class="area">{"".join(f"<li>{a}</li>" for a in AREA)}</ul></div></section>

<section><div class="wrap band"><div><h2>Пришлите фото — оценю работу</h2><p>Фото поверхности и примерная площадь. Отвечу, что нужно сделать и сколько это стоит.</p></div>
<div class="cta-row">{contact_html()}</div></div></section>
""", schema=[biz]))

urls.append(page("/uslugi/", "Услуги маляра в Сануре", "Все малярные услуги в Сануре: покраска, плесень, лак и масло для дерева, металл.",
                 f'<section><div class="wrap"><div class="sec-head"><h1>Услуги маляра в Сануре</h1><p class="lead">Выберите работу — на странице услуги этапы, материалы и частые вопросы.</p></div>{svc_cards()}</div></section>',
                 crumbs=[("Главная", "/"), ("Услуги", "")]))
for p in svcs:
    intro, steps = TXT[p["id"]]
    faq = faq_for(p["id"])
    faq_html = ('<section class="alt-bg"><div class="wrap"><div class="sec-head"><h2>Частые вопросы</h2></div>' + "".join(
        f"<details><summary>{E(q[0].upper() + q[1:])}?</summary><p>Зависит от состояния поверхности и материала. Пришлите фото — подскажу, что нужно именно в вашем случае.</p></details>"
        for q in faq) + "</div></section>") if faq else ""
    svc_schema = {"@context": "https://schema.org", "@type": "Service", "name": p["h1"], "areaServed": "Sanur, Bali",
                  "provider": {"@type": "HousePainter", "name": CFG["name"]}}
    crumbs_schema = {"@context": "https://schema.org", "@type": "BreadcrumbList", "itemListElement": [
        {"@type": "ListItem", "position": 1, "name": "Главная", "item": CFG["site_url"] + "/"},
        {"@type": "ListItem", "position": 2, "name": "Услуги", "item": CFG["site_url"] + "/uslugi/"},
        {"@type": "ListItem", "position": 3, "name": p["h1"]}]}
    others = [o for o in svcs if o["id"] != p["id"]][:3]
    urls.append(page(p["url"], p["title"], p["desc"], f"""
<div class="wrap page-hero"><div><span class="eyebrow">Санур · расчёт по фото</span><h1>{E(p["h1"])}</h1><p class="lead">{E(intro)}</p>
<div class="cta-row">{contact_html()}</div></div><div class="photo">{img(IMG[p["id"]], eager=True)}</div></div>
<section><div class="wrap split"><div><h2>Что входит в работу</h2><ol class="steps" style="grid-template-columns:1fr">{"".join(f"<li>{E(s)}</li>" for s in steps)}</ol></div>
<div><h2>Стоимость</h2><p>Цена зависит от площади, состояния поверхности и материалов. Точную смету даю после фото или осмотра — выезд в пределах Санура.</p>
<ul class="checks"><li>Смета до начала работ</li><li>Материалы под влажный климат</li><li>Уборка после работы</li></ul></div></div></section>
{faq_html}
<section><div class="wrap band"><div><h2>Нужна оценка?</h2><p>Пришлите фото поверхности и примерную площадь.</p></div><div class="cta-row">{contact_html()}</div></div></section>
<section><div class="wrap"><div class="sec-head"><h2>Другие услуги</h2></div>{svc_cards(others)}</div></section>
""", crumbs=[("Главная", "/"), ("Услуги", "/uslugi/"), (p["h1"], "")], schema=[svc_schema, crumbs_schema], og=IMG[p["id"]]))

for pid, body, pic in [("P090", "<p class=\"lead\">Цена складывается из площади, подготовки (очистка, обработка от плесени, шпаклёвка), материалов и доступа (высота, леса). Пришлите фото — пришлю расчёт.</p>", "roller"),
                       ("P091", "<p class=\"lead\">Здесь появятся фотографии реальных объектов в Сануре: до и после.</p>", "villa"),
                       ("P092", "<p class=\"lead\">Работаю только в Сануре. Напишите — отвечу и договоримся об осмотре.</p>", "bali")]:
    p = pages[pid]
    urls.append(page(p["url"], p["title"], p["desc"],
                     f'<div class="wrap page-hero"><div><h1>{E(p["h1"])}</h1>{body}<div class="cta-row">{contact_html()}</div></div><div class="photo">{img(pic, eager=True)}</div></div>',
                     crumbs=[("Главная", "/"), (p["h1"], "")], og=pic))

open(os.path.join(OUT, "sitemap.xml"), "w", encoding="utf-8").write(
    '<?xml version="1.0" encoding="UTF-8"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">' +
    "".join(f"<url><loc>{u}</loc></url>" for u in urls) + "</urlset>")
open(os.path.join(OUT, "robots.txt"), "w").write(f"User-agent: *\nAllow: /\nSitemap: {CFG['site_url']}/sitemap.xml\n")
open(os.path.join(OUT, "404.html"), "w", encoding="utf-8").write(f'<!doctype html><meta charset="utf-8"><title>Страница не найдена</title><p>Страница не найдена. <a href="{B}/">На главную</a></p>')
print("pages:", len(urls))
